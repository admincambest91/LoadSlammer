import json
import openpyxl
from openpyxl.cell.cell import MergedCell

class VOTFExtractor:
    def __init__(self, excel_path: str, sheet_name: str = "VOTF_Timing"):
        self.excel_path = excel_path
        self.sheet_name = sheet_name
        self.keywords = {
            "test case":  "Test_case",
            "psi mode":   "PSI_Mode",
            "idd set":    "IDDSet",
            "old vid":    "Old_VID",
            "target vid": "Target_VID",
        }
        self.results = {name: (None, None) for name in self.keywords.values()}
        self.workbook = None
        self.sheet = None

    def load_workbook(self):
        print(f"Loading workbook: {self.excel_path}")
        self.workbook = openpyxl.load_workbook(self.excel_path, data_only=True)
        self.sheet = self.workbook[self.sheet_name]
        print(f"Workbook loaded, using sheet: {self.sheet_name}")

    @staticmethod
    def get_merged_cell_value(sheet, row: int, col: int):
        """Return the value of a cell, resolving merged-cell references."""
        cell = sheet.cell(row=row, column=col)
        if not isinstance(cell, MergedCell):
            return cell.value
        for rng in sheet.merged_cells.ranges:
            if cell.coordinate in rng:
                top_left = sheet.cell(row=rng.min_row, column=rng.min_col)
                return top_left.value
        return None

    def find_headers(self, max_row=20, max_col=5):
        """Scan the top-left block for header keywords, store their positions."""
        for row in self.sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_col):
            for cell in row:
                val = cell.value
                if not val:
                    continue
                txt = str(val).lower()
                for substr, friendly in self.keywords.items():
                    if self.results[friendly] == (None, None) and substr in txt:
                        self.results[friendly] = (cell.row, cell.column)
            if all(pos != (None, None) for pos in self.results.values()):
                break

    def extract_column(self, header_pos, transform=lambda x: x):
        """Extract a list from a single column starting below header_pos."""
        row_h, col_h = header_pos
        data = []
        for row in self.sheet.iter_rows(min_row=row_h + 1, max_row=19, min_col=col_h, max_col=col_h):
            cell = row[0]
            if isinstance(cell, MergedCell):
                raw = self.get_merged_cell_value(self.sheet, cell.row, cell.column)
            else:
                raw = cell.value
            if raw is None:
                continue
            try:
                val = transform(raw)
            except Exception:
                continue
            if val is not None:
                data.append(val)
        return data

    def build_data(self):
        """Run the full extraction and build the JSON-serializable dict."""
        # 1. Find headers
        self.find_headers()

        # 2. Unpack header positions
        tc_pos = self.results["Test_case"]
        psi_pos = self.results["PSI_Mode"]
        idd_pos = self.results["IDDSet"]
        old_pos = self.results["Old_VID"]
        tgt_pos = self.results["Target_VID"]

        # 3. Extract columns
        test_cases = self.extract_column(tc_pos, transform=lambda v: abs(v) if isinstance(v, (int, float)) else v)
        psi_modes  = self.extract_column(psi_pos)
        idd_sets   = self.extract_column(idd_pos, transform=lambda v: abs(v) if isinstance(v, (int, float)) else v)
        old_vids   = self.extract_column(old_pos)
        tgt_vids   = self.extract_column(tgt_pos)

        # 4. Build mapping
        data = {
            str(tc): {
                "PSI_Mode":   psi_modes[i],
                "IDD_Set":    idd_sets[i],
                "Old_VID":    old_vids[i],
                "Target_VID": tgt_vids[i],
            }
            for i, tc in enumerate(test_cases)
        }
        return data

    def save_to_json(self, data: dict, output_path: str = "votf_parameter.json"):
        """Write the extracted data out to a JSON file."""
        with open(output_path, "w") as f:
            json.dump(data, f, indent=2)
        print(f"Data saved to {output_path}")

    def run(self, output_path: str = "votf_parameter.json"):
        """Full end-to-end execution: load, parse, and save."""
        self.load_workbook()
        data = self.build_data()
        self.save_to_json(data, output_path)


if __name__ == "__main__":
    CPU0_SOC_mainsheet="C:\\Users\\HPS Penang Tester\\Documents\\python\Load_slammer\\LoadSlammer_Testplan\\SP5_CPU_SVI3_VDDCRCPU0_VDDCRSOC_Analysis_V0_2.xlsm"

    extractor = VOTFExtractor(CPU0_SOC_mainsheet)
    extractor.run()
