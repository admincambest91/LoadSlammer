import openpyxl
from openpyxl.utils import column_index_from_string
from openpyxl.drawing.image import Image
from time import sleep
import json
import os

CPU0_SOC_mainsheet="C:\\Users\\HPS Penang Tester\\Documents\\python\Load_slammer\\LoadSlammer_Testplan\\SP5_CPU_SVI3_VDDCRCPU0_VDDCRSOC_Analysis_V0_2.xlsm"
CPU1_VDDIO_mainsheet="C:\\Users\\HPS Penang Tester\\Documents\\python\\Load_slammer\\LoadSlammer_Testplan\\SP5_CPU_SVI3_VDDCRCPU1_VDDIO_ Analysis_V0_2.xlsm"


class JSON_excel_extractor:
    def __init__(self, JSON_file,excel_path: str, sheet_name: str = "VOTF_Timing"):
        
        self.excel_path = excel_path
        self.sheet_name = sheet_name
        self.json=JSON_file
        self.workbook = None
        self.sheet = None

    def load_workbook(self):
        print(f"Loading workbook: {self.excel_path}")
        self.workbook = openpyxl.load_workbook(self.excel_path, data_only=True)
        self.sheet = self.workbook[self.sheet_name]
        print(f"Workbook loaded, using sheet: {self.sheet_name}")
        return self.workbook, self.sheet

    def load_json_file(self):
        with open(self.json, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return data
    

    def save_output_data(self,output_data, output_dir="Result", filename="Output_data.json"):
        """
        Save the given dict to a JSON file inside output_dir/filename.
        Creates output_dir if it doesn't exist, and overwrites the file each time.
        
        Args:
            output_data (dict):  The data to serialize.
            output_dir (str):    Folder to place the file in.
            filename (str):      JSON filename.
        
        Returns:
            str:  Full path to the written file.
        """
        # 1. Ensure the directory exists
        os.makedirs(output_dir, exist_ok=True)
        
        # 2. Build the full path
        self.full_path = os.path.join(output_dir, filename)
        
        # 3. Write (and overwrite) the JSON file
        with open(self.full_path, "w") as f:
            json.dump(output_data, f, indent=4)
        
        print(f"[INFO] Saved output data to: {self.full_path}")
        return self.full_path

    

if __name__ == "__main__":
    extractor=JSON_excel_extractor("CPU0_SOC_votf_parameter.json",CPU0_SOC_mainsheet)
    json_data=extractor.load_json_file()
    wb,sheet=extractor.load_workbook()
    sleep(1)
